from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from employees.models import Employee


def import_employees(file):
    workbook = load_workbook(file)

    sheet = workbook.active
    if sheet is None:
        raise ValueError("No active worksheet found.")

    worksheet: Worksheet = sheet

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        (
            employee_id,
            first_name,
            last_name,
            email,
            phone,
            department,
            designation,
            salary,
            joining_date,
        ) = row

        Employee.objects.update_or_create(
            employee_id=employee_id,
            defaults={
                "first_name": first_name,
                "last_name": last_name,
                "email": email,
                "phone": phone,
                "department": department,
                "designation": designation,
                "salary": salary,
                "joining_date": joining_date,
            },
        )